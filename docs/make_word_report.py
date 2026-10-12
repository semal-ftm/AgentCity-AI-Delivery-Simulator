"""Build docs/REPORT.docx from docs/REPORT.md (diagrams and result charts included).

    python docs/make_word_report.py

Needs python-docx (pip install python-docx). Run results/diagrams/*.png must exist
(they are rendered from the SVG diagrams).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS = Path(__file__).resolve().parent
ROOT = DOCS.parent
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
CHARTS = [("1_delivery_time.png", "Mean delivery time per agent (dashed line = deadline)"),
          ("2_cost.png", "Mean delivery cost per agent"),
          ("3_safety.png", "Mean safety per agent"),
          ("4_deadline_success.png", "Deliveries on time per agent"),
          ("5_reward.png", "Mean reward (overall performance) per agent"),
          ("6_learning_curve.png", "Learning agent: reward during training")]


def shade(cell, hex_color):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pr.append(shd)


def add_runs(par, text, size=None, bold_all=False):
    """Inline markdown: **bold**, *italic*, `code`."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2]); run.bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1]); run.font.name = "Consolas"
        elif part.startswith("*"):
            run = par.add_run(part[1:-1]); run.italic = True
        else:
            run = par.add_run(part)
        if bold_all:
            run.bold = True
        if size:
            run.font.size = Pt(size)


def add_table(doc, rows, cover=False):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if cover:
        cells = [r for r in cells if any(r)]
    n = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(cells):
        for j in range(n):
            cell = t.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            add_runs(par, r[j] if j < len(r) else "", size=11 if cover else 9.5, bold_all=(i == 0 and not cover))
            if i == 0 and not cover:
                shade(cell, "DCE6F1")
            if cover and j == 0:
                shade(cell, "EEF2F7")
    doc.add_paragraph()


def add_picture(doc, path, caption=None, width=16):
    if not path.exists():
        doc.add_paragraph(f"[missing image: {path.name}]")
        return
    doc.add_picture(str(path), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption); r.italic = True; r.font.size = Pt(9)


def build():
    md = (DOCS / "REPORT.md").read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for lvl, size in ((1, 16), (2, 14), (3, 12)):
        h = doc.styles[f"Heading {lvl}"]; h.font.color.rgb = NAVY; h.font.size = Pt(size)

    i, first_table_done = 0, False
    while i < len(md):
        line = md[i]
        if line.startswith("# "):                                   # title page
            for _ in range(4):
                doc.add_paragraph()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:]); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = NAVY
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("Seven AI agent architectures for one food-delivery scenario"); r.italic = True; r.font.size = Pt(13)
            doc.add_paragraph()
        elif line.startswith("|"):
            rows = []
            while i < len(md) and md[i].startswith("|"):
                rows.append(md[i]); i += 1
            add_table(doc, rows, cover=not first_table_done)
            first_table_done = True
            continue
        elif line.strip() == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif m := re.match(r"!\[(.*?)\]\((.*?)\)", line.strip()):
            path = (DOCS / m.group(2)).resolve().with_suffix(".png")
            add_picture(doc, path, f"Figure: {m.group(1)} agent architecture")
        elif line.startswith("```"):
            i += 1
            p = doc.add_paragraph()
            while i < len(md) and not md[i].startswith("```"):
                r = p.add_run(md[i] + "\n"); r.font.name = "Consolas"; r.font.size = Pt(9.5)
                i += 1
        elif line.startswith("> "):
            p = doc.add_paragraph(); add_runs(p, line[2:]); p.runs[0].italic = True
        elif re.match(r"^\s*- ", line):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, re.sub(r"^\s*- ", "", line))
        elif line.startswith("Charts:"):
            doc.add_paragraph("Charts (generated by run_experiments.py):")
            for f, cap in CHARTS:
                add_picture(doc, ROOT / "results" / "figures" / f, cap, width=15.5)
        elif line.strip():
            p = doc.add_paragraph(); add_runs(p, line.strip())
        i += 1
    out = DOCS / "REPORT.docx"
    doc.save(out)
    print("written", out)


if __name__ == "__main__":
    build()
